filename = '/Users/xiaobo/Desktop/test.xlsx'

from openpyxl import load_workbook
from openpyxl.chart import (
    ScatterChart,
    Reference,
    Series,
)

wb = load_workbook(filename)
ws = wb.active

chart = ScatterChart()

chart.style = 13
chart.x_axis.title = 'Time'
chart.y_axis.title = 'RFU'

xvalues = Reference(ws, min_col=1, min_row=2, max_row=21)     #starting col including x xalue, starting data row and finishing data row
for i in range(2,9):                                          #cols including y values, noted final number is n+1 which is python range feature
    values = Reference(ws, min_col=i, min_row=1, max_row=21)  #rows of everything including first title row
    series = Series(values, xvalues, title_from_data=True)
    chart.series.append(series)


ws.add_chart(chart, "D10")

wb.save(filename)
